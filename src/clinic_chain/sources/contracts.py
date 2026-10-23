"""Source contracts (`docs/contracts/*.yaml`): read output tables and check them against schema."""

import csv
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import yaml
from openpyxl import load_workbook

DATETIME = "%Y-%m-%d %H:%M:%S"


def load(contracts_dir: Path) -> dict[str, dict]:
    """All tables keyed `system.table`."""
    tables = {}
    for path in sorted(contracts_dir.glob("*.yaml")):
        doc = yaml.safe_load(path.read_text())
        for name, spec in doc["tables"].items():
            tables[f"{doc['system']}.{name}"] = spec
    return tables


def read(root: Path, spec: dict) -> tuple[list[str], list[list]]:
    """Header and rows of one table; CSV empty fields come back as None."""
    file, _, sheet = spec["file"].partition("#")
    if sheet:
        ws = load_workbook(root / file, read_only=True)[sheet]
        header, *rows = ([c for c in row] for row in ws.iter_rows(values_only=True))
        return header, rows
    with (root / file).open(newline="", encoding="utf-8") as f:
        header, *rows = csv.reader(f)
    return header, [[v if v != "" else None for v in row] for row in rows]


def _type_ok(value, kind: str) -> bool:
    try:
        if kind == "string":
            return isinstance(value, str)
        if kind == "int":
            return isinstance(value, int) or str(int(value)) == value
        if kind == "decimal":
            Decimal(str(value))
            return True
        if kind == "date":
            if isinstance(value, datetime):
                return value.time() == datetime.min.time()
            return isinstance(value, date) or bool(date.fromisoformat(value)) and len(value) == 10
        if kind == "datetime":
            return isinstance(value, datetime) or bool(datetime.strptime(value, DATETIME))
    except (ValueError, TypeError, InvalidOperation):
        return False
    raise ValueError(f"unknown contract type {kind}")


def validate(root: Path, name: str, spec: dict) -> list[str]:
    """Contract violations of one table; empty when it conforms."""
    header, rows = read(root, spec)
    columns = spec["columns"]
    expected = [c["name"] for c in columns]
    if header != expected:
        return [f"{name}: header {header} != contract {expected}"]
    errors = []
    for i, row in enumerate(rows, 2):
        for col, value in zip(columns, row, strict=True):
            if value is None:
                if not col.get("nullable"):
                    errors.append(f"{name} row {i}: {col['name']} is null")
            elif not _type_ok(value, col["type"]):
                errors.append(f"{name} row {i}: {col['name']}={value!r} is not {col['type']}")
        if len(errors) > 20:
            return errors
    key = [expected.index(k) for k in spec["primary_key"]]
    keys = [tuple(row[i] for i in key) for row in rows]
    if len(set(keys)) != len(keys):
        errors.append(f"{name}: primary key {spec['primary_key']} is not unique")
    return errors
